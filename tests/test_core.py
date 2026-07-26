from __future__ import annotations

import os
import tempfile
import unittest
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PIL import Image
from PySide6.QtWidgets import QApplication
from pptx import Presentation

from image_utils import build_asset_path
from main import PPTCloneApp, classify_import_paths, parse_inpaint_api_urls
from persistence import atomic_write_json, load_json_with_backup
from ppt_export import PPTExporter


class InpaintUrlTests(unittest.TestCase):
    def test_parse_urls_splits_deduplicates_and_preserves_order(self):
        value = "http://127.0.0.1:8080/a; http://localhost:8081/b\nhttp://127.0.0.1:8080/a"
        self.assertEqual(
            parse_inpaint_api_urls(value),
            ["http://127.0.0.1:8080/a", "http://localhost:8081/b"],
        )

    def test_parse_urls_accepts_sequence(self):
        self.assertEqual(
            parse_inpaint_api_urls(
                [
                    "http://localhost/a;http://localhost/b",
                    "",
                    "http://localhost/a",
                ]
            ),
            ["http://localhost/a", "http://localhost/b"],
        )


class ImportClassificationTests(unittest.TestCase):
    def test_mixed_paths_are_classified_and_deduplicated(self):
        images, pdfs, unsupported = classify_import_paths(
            ["scan.PNG", "deck.pdf", "scan.PNG", "notes.txt"]
        )
        self.assertEqual([os.path.basename(path) for path in images], ["scan.PNG"])
        self.assertEqual([os.path.basename(path) for path in pdfs], ["deck.pdf"])
        self.assertEqual([os.path.basename(path) for path in unsupported], ["notes.txt"])


class PersistenceTests(unittest.TestCase):
    def test_atomic_json_write_keeps_and_recovers_previous_version(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            settings_path = os.path.join(temp_dir, "settings.json")
            atomic_write_json(settings_path, {"generation": 1, "label": "中文"})
            atomic_write_json(settings_path, {"generation": 2})

            with open(settings_path, "w", encoding="utf-8") as handle:
                handle.write("{broken")

            recovered, used_backup = load_json_with_backup(settings_path)
            self.assertTrue(used_backup)
            self.assertEqual(recovered, {"generation": 1, "label": "中文"})


class AssetPathTests(unittest.TestCase):
    def test_asset_path_is_stable_and_safe(self):
        first = build_asset_path("cache", "slide", r"C:\input\课程 封面.png", suffix="preview", ext="jpg")
        second = build_asset_path("cache", "slide", r"C:\input\课程 封面.png", suffix="preview", ext="jpg")
        self.assertEqual(first, second)
        self.assertTrue(first.endswith("_preview.jpg"))
        self.assertNotIn(" ", os.path.basename(first))


class PowerPointExportTests(unittest.TestCase):
    def test_minimal_pptx_can_be_reopened(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = os.path.join(temp_dir, "slide.png")
            output_path = os.path.join(temp_dir, "result.pptx")
            Image.new("RGB", (640, 360), "white").save(image_path)

            exporter = PPTExporter()
            exporter.add_image_with_text_boxes(
                image_path,
                [{"text": "Editable text", "rect": [40, 40, 240, 60], "confidence": 0.99}],
            )
            exporter.save(output_path)

            presentation = Presentation(output_path)
            self.assertEqual(len(presentation.slides), 1)
            self.assertTrue(any(shape.has_text_frame and "Editable text" in shape.text for shape in presentation.slides[0].shapes))

    def test_failed_export_preserves_existing_file(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "existing.pptx")
            original = b"keep this presentation"
            with open(output_path, "wb") as handle:
                handle.write(original)

            exporter = PPTExporter()
            with mock.patch.object(exporter.prs, "save", side_effect=OSError("disk full")):
                self.assertFalse(exporter.save(output_path))

            with open(output_path, "rb") as handle:
                self.assertEqual(handle.read(), original)
            self.assertEqual(os.listdir(temp_dir), ["existing.pptx"])


class ApplicationSmokeTests(unittest.TestCase):
    def test_window_can_start_and_close_offscreen(self):
        app = QApplication.instance() or QApplication([])
        window = PPTCloneApp()
        try:
            self.assertEqual(window.windowTitle(), "PowerOCR 演示" if window.ui_lang == "zh" else "PowerOCR Presentation")
        finally:
            window.close()
            app.processEvents()


if __name__ == "__main__":
    unittest.main()
